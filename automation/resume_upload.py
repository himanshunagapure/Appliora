"""
Prepare ATS-friendly resumes for upload.

Canonical apply artifact is a single-column Calibri DOCX (no tables, headers,
text boxes, or graphics). Markdown is kept as a human-review sibling.
"""

from __future__ import annotations

import re
from pathlib import Path

from utils.logger import get_logger

logger = get_logger("jobapplypro.resume_upload")

UPLOADABLE_SUFFIXES = {".pdf", ".doc", ".docx"}
ATS_FONT = "Calibri"

# Map common LLM headings onto names ATS parsers actually look for.
_SECTION_ALIASES = {
    "summary": "SUMMARY",
    "professional summary": "SUMMARY",
    "profile": "SUMMARY",
    "objective": "SUMMARY",
    "about": "SUMMARY",
    "experience": "EXPERIENCE",
    "work experience": "EXPERIENCE",
    "professional experience": "EXPERIENCE",
    "employment": "EXPERIENCE",
    "work history": "EXPERIENCE",
    "skills": "SKILLS",
    "technical skills": "SKILLS",
    "core skills": "SKILLS",
    "key skills": "SKILLS",
    "projects": "PROJECTS",
    "selected projects": "PROJECTS",
    "personal projects": "PROJECTS",
    "education": "EDUCATION",
    "academics": "EDUCATION",
    "academic": "EDUCATION",
    "certifications": "CERTIFICATIONS",
    "certificates": "CERTIFICATIONS",
    "licenses": "CERTIFICATIONS",
}


def ensure_uploadable_resume(path: str | Path, *, output_dir: Path | None = None) -> Path:
    """
    Return a path LinkedIn/ATS file inputs will accept.

    - PDF/DOC/DOCX are returned as-is.
    - Markdown/text are converted to a sibling ATS-friendly .docx
      (cached if newer than source). PDF sibling is also written when converting.
    """
    src = Path(path)
    if not src.exists():
        raise FileNotFoundError(f"Resume not found: {src}")

    suffix = src.suffix.lower()
    if suffix in UPLOADABLE_SUFFIXES:
        return src

    if suffix not in {".md", ".markdown", ".txt", ""}:
        raise ValueError(
            f"Unsupported resume format {suffix!r}; need pdf/doc/docx or md/txt"
        )

    out_dir = output_dir or src.parent
    out_dir.mkdir(parents=True, exist_ok=True)
    dest = out_dir / f"{src.stem}.docx"
    pdf_dest = out_dir / f"{src.stem}.pdf"
    body = src.read_text(encoding="utf-8")

    if not (dest.exists() and dest.stat().st_mtime >= src.stat().st_mtime):
        markdown_to_docx(body, dest)
        logger.info("Converted resume %s → ATS-friendly %s", src.name, dest)
    if not (pdf_dest.exists() and pdf_dest.stat().st_mtime >= src.stat().st_mtime):
        try:
            markdown_to_pdf(body, pdf_dest)
        except Exception as exc:
            logger.warning("PDF conversion skipped for %s: %s", src.name, exc)
            try:
                bad = pdf_dest.with_suffix(pdf_dest.suffix + ".tmp")
                bad.unlink(missing_ok=True)
            except OSError:
                pass
            if pdf_dest.exists() and pdf_dest.stat().st_size < 100:
                try:
                    pdf_dest.unlink()
                except OSError:
                    pass
    styled_dest = out_dir / f"{src.stem}.styled.pdf"
    if not (styled_dest.exists() and styled_dest.stat().st_mtime >= src.stat().st_mtime):
        try:
            markdown_to_styled_pdf(body, styled_dest)
        except Exception as extra:
            logger.warning("Styled PDF skipped for %s: %s", src.name, extra)
    return dest


def read_resume_text(path: str | Path, *, max_chars: int | None = None) -> str:
    """
    Plain text for LLM prompts / review excerpts.

    Prefer the markdown sibling of a .docx/.pdf so resume_facts stay verbatim.
    """
    p = Path(path)
    suffix = p.suffix.lower()
    text = ""
    if suffix in {".md", ".markdown", ".txt", ""}:
        text = p.read_text(encoding="utf-8")
    elif suffix in {".docx", ".pdf"}:
        sibling = p.with_suffix(".md")
        if not sibling.exists():
            sibling = p.with_suffix(".txt")
        if sibling.exists():
            text = sibling.read_text(encoding="utf-8")
        elif suffix == ".docx":
            try:
                text = docx_to_text(p)
            except Exception as exc:
                logger.warning("Could not extract text from %s: %s", p, exc)
                text = ""
        else:
            try:
                text = pdf_to_text(p)
            except Exception as exc:
                logger.warning("Could not extract text from %s: %s", p, exc)
                text = ""
    elif suffix in {".doc"}:
        raise ValueError(f"Cannot extract text from legacy .doc: {p}")
    else:
        text = p.read_text(encoding="utf-8", errors="replace")
    if max_chars is not None and len(text) > max_chars:
        return text[:max_chars] + "\n\n[truncated]"
    return text


def docx_to_text(path: str | Path) -> str:
    from docx import Document

    doc = Document(str(path))
    return "\n".join(p.text for p in doc.paragraphs if p.text.strip())


def pdf_to_text(path: str | Path) -> str:
    from pypdf import PdfReader

    reader = PdfReader(str(path))
    parts: list[str] = []
    for page in reader.pages:
        parts.append(page.extract_text() or "")
    return "\n".join(parts).strip()


def iter_resume_blocks(markdown: str) -> list[tuple[str, str]]:
    """
    Parse tailored markdown into typed blocks for DOCX/PDF renderers.

    Block kinds: name | contact | section | job | bullet | body
    """
    blocks: list[tuple[str, str]] = []
    lines = (markdown or "").replace("\r\n", "\n").split("\n")
    saw_name = False
    i = 0
    while i < len(lines):
        stripped = lines[i].strip()
        if not stripped:
            i += 1
            continue
        if re.fullmatch(r"-{3,}|_{3,}|\*{3,}", stripped):
            i += 1
            continue

        heading_match = re.match(r"^(#{1,3})\s+(.*)$", stripped)
        if heading_match:
            level = len(heading_match.group(1))
            title = _strip_md_inline(heading_match.group(2))
            if level == 1 and not saw_name:
                blocks.append(("name", title))
                saw_name = True
            elif level in {1, 2}:
                blocks.append(("section", _normalize_section(title)))
            else:
                blocks.append(("job", title))
            i += 1
            continue

        if i + 1 < len(lines) and re.fullmatch(r"-{3,}", lines[i + 1].strip()):
            blocks.append(("section", _normalize_section(_strip_md_inline(stripped))))
            i += 2
            continue

        bullet_match = re.match(r"^(\s*)(?:[\*\-\u2022]|\d+[.)])\s+(.*)$", lines[i].rstrip())
        if bullet_match:
            blocks.append(("bullet", _strip_md_inline(bullet_match.group(2))))
            i += 1
            continue

        text = _strip_md_inline(stripped)
        if not saw_name and _looks_like_name(text):
            blocks.append(("name", text))
            saw_name = True
        elif _looks_like_contact(text):
            blocks.append(("contact", text))
        else:
            blocks.append(("body", text))
        i += 1
    return blocks


def markdown_to_docx(markdown: str, dest: Path) -> Path:
    """
    Render markdown as an ATS-friendly single-column DOCX.

    ATS rules:
      - Calibri, black text, 11pt body
      - No tables, columns, headers, footers, images, or text boxes
      - Standard ALL-CAPS section headings
      - Plain hyphen bullets (not Word list XML)
    """
    from docx import Document
    from docx.shared import Inches

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()
    _configure_ats_styles(doc)
    _clear_default_empty_paragraph(doc)

    section = doc.sections[0]
    section.top_margin = Inches(0.6)
    section.bottom_margin = Inches(0.6)
    section.left_margin = Inches(0.7)
    section.right_margin = Inches(0.7)
    section.header_distance = Inches(0.0)
    section.footer_distance = Inches(0.0)
    for hf in (section.header, section.footer):
        hf.is_linked_to_previous = False
        for para in hf.paragraphs:
            para.text = ""

    blocks = iter_resume_blocks(markdown)
    if not blocks:
        _add_body(doc, "")
    for kind, text in blocks:
        if kind == "name":
            _add_name(doc, text)
        elif kind == "contact":
            _add_contact(doc, text)
        elif kind == "section":
            _add_section_heading_label(doc, text)
        elif kind == "job":
            _add_job_line(doc, text)
        elif kind == "bullet":
            _add_bullet(doc, text)
        else:
            _add_body(doc, text)

    doc.save(str(dest))
    logger.debug("Wrote ATS-friendly DOCX %s (%d paragraphs)", dest, len(doc.paragraphs))
    return dest


def markdown_to_pdf(markdown: str, dest: Path) -> Path:
    """
    Render markdown as an ATS-friendly text PDF (selectable text, not a scan).

    Uses Helvetica (standard PDF fonts) — single column, black text, no tables
    or graphics. Companion to the Calibri DOCX upload artifact.
    Writes atomically (temp file → rename) so a failed run cannot leave a
    corrupt .pdf that would skip regeneration on the next convert.
    """
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_suffix(dest.suffix + ".tmp")

    doc = SimpleDocTemplate(
        str(tmp),
        pagesize=letter,
        leftMargin=0.7 * inch,
        rightMargin=0.7 * inch,
        topMargin=0.6 * inch,
        bottomMargin=0.6 * inch,
        title="Resume",
        author="JobApplyPro",
    )
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "ATSName",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=16,
        leading=20,
        textColor="black",
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "ATSContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        textColor="black",
        spaceAfter=10,
    )
    section_style = ParagraphStyle(
        "ATSSection",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=12,
        leading=15,
        textColor="black",
        spaceBefore=10,
        spaceAfter=3,
    )
    job_style = ParagraphStyle(
        "ATSJob",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor="black",
        spaceBefore=6,
        spaceAfter=1,
    )
    body_style = ParagraphStyle(
        "ATSBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=11,
        leading=14,
        textColor="black",
        spaceAfter=4,
    )
    bullet_style = ParagraphStyle(
        "ATSBullet",
        parent=body_style,
        leftIndent=12,
        spaceAfter=2,
    )

    story: list = []
    for kind, text in iter_resume_blocks(markdown):
        safe = _pdf_escape(text)
        if kind == "name":
            story.append(Paragraph(safe, name_style))
        elif kind == "contact":
            story.append(Paragraph(safe, contact_style))
        elif kind == "section":
            story.append(Paragraph(safe, section_style))
        elif kind == "job":
            story.append(Paragraph(safe, job_style))
        elif kind == "bullet":
            story.append(Paragraph(f"- {safe}", bullet_style))
        else:
            story.append(Paragraph(safe, body_style))
    if not story:
        story.append(Spacer(1, 1))

    try:
        doc.build(story)
        tmp.replace(dest)
    except Exception:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        raise
    logger.debug("Wrote ATS-friendly PDF %s", dest)
    return dest


NAVY = "#1B365D"
GOLD = "#C4A35A"
INK = "#1A1A1A"
MUTED = "#5A6570"
RULE = "#D6DCE3"


def markdown_to_styled_pdf(markdown: str, dest: Path) -> Path:
    """
    Visual/layout resume PDF (Phase 11). Not for ATS upload.

    Navy header band, gold rules, colored section titles. Companion to the
    Helvetica ATS PDF — written as `*.styled.pdf` so parsers never pick it up.
    Atomic write (temp → rename).
    """
    from reportlab.lib.colors import HexColor, white
    from reportlab.lib.enums import TA_CENTER, TA_LEFT
    from reportlab.lib.pagesizes import letter
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import inch
    from reportlab.platypus import (
        HRFlowable,
        KeepTogether,
        Paragraph,
        SimpleDocTemplate,
        Spacer,
        Table,
        TableStyle,
    )

    dest = Path(dest)
    dest.parent.mkdir(parents=True, exist_ok=True)
    tmp = dest.with_suffix(dest.suffix + ".tmp")

    navy = HexColor(NAVY)
    gold = HexColor(GOLD)
    ink = HexColor(INK)
    muted = HexColor(MUTED)
    rule = HexColor(RULE)

    doc = SimpleDocTemplate(
        str(tmp),
        pagesize=letter,
        leftMargin=0.55 * inch,
        rightMargin=0.55 * inch,
        topMargin=0.45 * inch,
        bottomMargin=0.5 * inch,
        title="Resume",
        author="JobApplyPro",
    )
    styles = getSampleStyleSheet()
    name_style = ParagraphStyle(
        "StyledName",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=20,
        leading=24,
        textColor=white,
        alignment=TA_CENTER,
        spaceAfter=2,
    )
    contact_style = ParagraphStyle(
        "StyledContact",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=9,
        leading=12,
        textColor=white,
        alignment=TA_CENTER,
        spaceAfter=0,
    )
    section_style = ParagraphStyle(
        "StyledSection",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=11,
        leading=14,
        textColor=navy,
        spaceBefore=8,
        spaceAfter=2,
        alignment=TA_LEFT,
    )
    job_style = ParagraphStyle(
        "StyledJob",
        parent=styles["Normal"],
        fontName="Helvetica-Bold",
        fontSize=10,
        leading=13,
        textColor=ink,
        spaceBefore=6,
        spaceAfter=1,
    )
    body_style = ParagraphStyle(
        "StyledBody",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=10,
        leading=13,
        textColor=ink,
        spaceAfter=3,
    )
    bullet_style = ParagraphStyle(
        "StyledBullet",
        parent=body_style,
        leftIndent=14,
        bulletIndent=0,
        spaceAfter=2,
    )
    footer_style = ParagraphStyle(
        "StyledFooter",
        parent=styles["Normal"],
        fontName="Helvetica",
        fontSize=8,
        textColor=muted,
        alignment=TA_CENTER,
    )

    story: list = []
    header_bits: list[tuple[str, str]] = []
    rest: list[tuple[str, str]] = []
    for kind, text in iter_resume_blocks(markdown):
        if kind in {"name", "contact"} and not rest:
            header_bits.append((kind, text))
        else:
            rest.append((kind, text))

    name_html = ""
    contact_html = ""
    for kind, text in header_bits:
        if kind == "name" and not name_html:
            name_html = _pdf_escape(text)
        else:
            contact_html = _pdf_escape(text)
    header_flow = []
    if name_html:
        header_flow.append(Paragraph(name_html, name_style))
    if contact_html:
        header_flow.append(Spacer(1, 4))
        header_flow.append(Paragraph(contact_html, contact_style))
    if header_flow:
        inner = Table([[header_flow]], colWidths=[7.4 * inch])
        inner.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, -1), navy),
                    ("LEFTPADDING", (0, 0), (-1, -1), 14),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 14),
                    ("TOPPADDING", (0, 0), (-1, -1), 14),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 14),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                ]
            )
        )
        story.append(inner)
        story.append(HRFlowable(width="100%", thickness=3, color=gold, spaceAfter=8, spaceBefore=0))

    for kind, text in rest:
        safe = _pdf_escape(text)
        if kind == "name":
            story.append(Paragraph(safe, job_style))
        elif kind == "contact":
            story.append(Paragraph(safe, body_style))
        elif kind == "section":
            block = [
                Paragraph(safe, section_style),
                HRFlowable(width="100%", thickness=0.8, color=gold, spaceAfter=4, spaceBefore=0),
            ]
            story.append(KeepTogether(block))
        elif kind == "job":
            story.append(Paragraph(safe, job_style))
        elif kind == "bullet":
            story.append(Paragraph(f"- {safe}", bullet_style))
        else:
            story.append(Paragraph(safe, body_style))

    if not story:
        story.append(Spacer(1, 1))
    story.append(Spacer(1, 16))
    story.append(HRFlowable(width="100%", thickness=0.4, color=rule, spaceBefore=4, spaceAfter=6))
    story.append(Paragraph("Designed layout - JobApplyPro", footer_style))

    try:
        doc.build(story)
        tmp.replace(dest)
    except Exception:
        try:
            tmp.unlink(missing_ok=True)
        except OSError:
            pass
        raise
    logger.debug("Wrote styled resume PDF %s", dest)
    return dest


def _pdf_escape(text: str) -> str:
    return (
        (text or "")
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
    )


def _clear_default_empty_paragraph(doc) -> None:
    """python-docx starts every Document with a blank paragraph — drop it."""
    if len(doc.paragraphs) == 1 and not doc.paragraphs[0].text.strip():
        el = doc.paragraphs[0]._element
        parent = el.getparent()
        if parent is not None:
            parent.remove(el)


def _configure_ats_styles(doc) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    def _font(style, name: str, size_pt: float, *, bold: bool = False) -> None:
        style.font.name = name
        style.font.size = Pt(size_pt)
        style.font.bold = bold
        style.font.color.rgb = RGBColor(0, 0, 0)
        rpr = style.element.get_or_add_rPr()
        rfonts = rpr.find(qn("w:rFonts"))
        if rfonts is None:
            from docx.oxml import OxmlElement

            rfonts = OxmlElement("w:rFonts")
            rpr.insert(0, rfonts)
        rfonts.set(qn("w:ascii"), name)
        rfonts.set(qn("w:hAnsi"), name)
        rfonts.set(qn("w:eastAsia"), name)
        rfonts.set(qn("w:cs"), name)

    _font(doc.styles["Normal"], ATS_FONT, 11)
    for heading_name, size in (("Heading 1", 16), ("Heading 2", 12)):
        try:
            _font(doc.styles[heading_name], ATS_FONT, size, bold=True)
        except KeyError:
            pass
    try:
        _font(doc.styles["List Bullet"], ATS_FONT, 11)
    except KeyError:
        pass


def _set_run_font(run, *, size_pt: float, bold: bool = False) -> None:
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor

    run.font.name = ATS_FONT
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.font.color.rgb = RGBColor(0, 0, 0)
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        from docx.oxml import OxmlElement

        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("w:ascii"), ATS_FONT)
    rfonts.set(qn("w:hAnsi"), ATS_FONT)
    rfonts.set(qn("w:eastAsia"), ATS_FONT)
    rfonts.set(qn("w:cs"), ATS_FONT)


def _space(paragraph, *, before: float = 0, after: float = 4) -> None:
    from docx.shared import Pt

    pf = paragraph.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = 1.08


def _add_name(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size_pt=16, bold=True)
    _space(p, before=0, after=2)


def _add_contact(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size_pt=10, bold=False)
    _space(p, before=0, after=8)


def _add_section_heading(doc, text: str) -> None:
    _add_section_heading_label(doc, _normalize_section(text))


def _add_section_heading_label(doc, label: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(label)
    _set_run_font(run, size_pt=12, bold=True)
    _space(p, before=10, after=2)


def _add_job_line(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size_pt=11, bold=True)
    _space(p, before=6, after=1)


def _add_bullet(doc, text: str) -> None:
    # Plain hyphen on a Normal paragraph — Word "List Bullet" numbering XML is
    # often skipped by ATS parsers.
    p = doc.add_paragraph()
    run = p.add_run(f"- {text}")
    _set_run_font(run, size_pt=11)
    _space(p, before=0, after=1)


def _add_body(doc, text: str) -> None:
    p = doc.add_paragraph()
    run = p.add_run(text)
    _set_run_font(run, size_pt=11)
    _space(p, before=0, after=4)


def _normalize_section(text: str) -> str:
    key = re.sub(r"[^a-z0-9 ]+", " ", text.strip().lower())
    key = re.sub(r"\s+", " ", key).strip()
    return _SECTION_ALIASES.get(key, text.strip().upper())


def _looks_like_name(text: str) -> bool:
    if "|" in text or "@" in text:
        return False
    words = [w for w in text.split() if w]
    if not (1 <= len(words) <= 4):
        return False
    # Avoid treating "Software Engineer" / title-case skill dumps as a name.
    lowered = {w.lower().strip(",.") for w in words}
    if lowered & {"engineer", "developer", "summary", "experience", "skills"}:
        return False
    return all(w[:1].isupper() for w in words)


_PHONE_RE = re.compile(
    r"(?:\+\d{1,3}[\s.-]*)?(?:\(?\d{2,4}\)?[\s.-]){2,}\d{3,4}"
)


def _looks_like_contact(text: str) -> bool:
    """Contact line only — not year ranges like 2019-2023 or GitHub in a bullet."""
    lower = text.lower()
    if "@" in text:
        return True
    has_link = "linkedin.com" in lower or "github.com" in lower
    has_phone = bool(_PHONE_RE.search(text))
    if "|" in text and (has_link or has_phone):
        return True
    if has_phone and len(text) <= 80:
        return True
    return False


def _strip_md_inline(text: str) -> str:
    """Remove markdown emphasis / links so ATS sees the real characters."""
    s = text
    s = re.sub(r"\[([^\]]+)\]\([^)]+\)", r"\1", s)
    s = re.sub(r"https?://\S+", lambda m: m.group(0).rstrip(").,;"), s)
    s = re.sub(r"\*\*([^*]+)\*\*", r"\1", s)
    s = re.sub(r"__([^_]+)__", r"\1", s)
    s = re.sub(r"\*([^*]+)\*", r"\1", s)
    s = re.sub(r"`([^`]+)`", r"\1", s)
    return s.strip()
